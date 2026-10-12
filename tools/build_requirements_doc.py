from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


PRESET = {
    "name": "standard_business_brief",
    "page_width": 8.5,
    "page_height": 11.0,
    "margin": 1.0,
    "header_distance": 0.492,
    "footer_distance": 0.492,
    "content_width_dxa": 9360,
    "table_indent_dxa": 120,
    "body_font": "Calibri",
    "body_east_asia": "微软雅黑",
    "body_size": 11,
    "body_after": 6,
    "body_line": 1.10,
    "h1_size": 16,
    "h1_before": 16,
    "h1_after": 8,
    "h2_size": 13,
    "h2_before": 12,
    "h2_after": 6,
    "h3_size": 12,
    "h3_before": 8,
    "h3_after": 4,
    "heading_blue": "2E74B5",
    "heading_dark_blue": "1F4D78",
    "ink": "24231F",
    "muted": "6F6B63",
    "table_fill": "F2F4F7",
    "table_alt_fill": "F7F8FA",
    "table_border": "D8DEE7",
    "callout_fill": "F4F6F9",
    "accent": "8A3A32",
}


CHAPTERS = {
    1: "第一章　项目基线与核心决策",
    5: "第二章　产品边界与系统上下文",
    14: "第三章　用户流程、状态与功能",
    21: "第四章　内容体系、推荐规则与AI",
    33: "第五章　交互原型与视觉规范",
    42: "第六章　技术架构与数据设计",
    60: "第七章　MVP实施路线与交付物",
    68: "第八章　质量、风险与最终验收",
    83: "第九章　当前落地状态与决策任务",
}


SKIP_SECTIONS = {13, 20, 32, 41, 59, 67, 82}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_table_borders(table, color: str = "D8DEE7", size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    if sum(widths_dxa) != PRESET["content_width_dxa"]:
        raise ValueError(f"Table widths must total 9360 DXA, got {sum(widths_dxa)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PRESET["content_width_dxa"]))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")


def table_widths(column_count: int, rows: list[list[str]]) -> list[int]:
    long_identifier_pattern = re.compile(r"[A-Za-z][A-Za-z0-9_]{15,}")
    if column_count == 2:
        first_max = max(len(row[0]) if row else 0 for row in rows)
        if first_max <= 12:
            return [1700, 7660]
        return [2700, 6660]
    if column_count == 3:
        if any(
            row and long_identifier_pattern.fullmatch(row[0]) for row in rows[1:]
        ):
            return [2300, 3530, 3530]
        return [1750, 3805, 3805]
    if column_count == 4:
        identifier_pattern = re.compile(r"[A-Z][A-Z0-9_]{12,}")
        if any(row and identifier_pattern.fullmatch(row[0]) for row in rows[1:]):
            return [2200, 2200, 2480, 2480]
        return [1350, 2300, 2855, 2855]
    if column_count == 5:
        return [1100, 1850, 1850, 2280, 2280]
    if column_count == 6:
        return [1200, 1550, 1550, 1550, 1755, 1755]
    base = PRESET["content_width_dxa"] // column_count
    widths = [base] * column_count
    widths[-1] += PRESET["content_width_dxa"] - sum(widths)
    return widths


def set_run_font(
    run,
    *,
    ascii_name: str = "Calibri",
    east_asia: str = "微软雅黑",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = ascii_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), ascii_name)
    r_fonts.set(qn("w:hAnsi"), ascii_name)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, ascii_name: str, east_asia: str, size: float, color: str | None = None) -> None:
    style.font.name = ascii_name
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), ascii_name)
    r_fonts.set(qn("w:hAnsi"), ascii_name)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, PRESET["body_font"], PRESET["body_east_asia"], PRESET["body_size"], PRESET["ink"])
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(PRESET["body_after"])
    normal.paragraph_format.line_spacing = PRESET["body_line"]

    for name, size, before, after, color in (
        ("Heading 1", PRESET["h1_size"], PRESET["h1_before"], PRESET["h1_after"], PRESET["heading_blue"]),
        ("Heading 2", PRESET["h2_size"], PRESET["h2_before"], PRESET["h2_after"], PRESET["heading_blue"]),
        ("Heading 3", PRESET["h3_size"], PRESET["h3_before"], PRESET["h3_after"], PRESET["heading_dark_blue"]),
    ):
        style = styles[name]
        set_style_font(style, "Calibri", "微软雅黑", size, color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    set_style_font(code_style, "Consolas", "等线", 8.5, "3E4550")
    code_style.paragraph_format.left_indent = Inches(0.14)
    code_style.paragraph_format.right_indent = Inches(0.10)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(6)
    code_style.paragraph_format.line_spacing = 1.08

    if "Caption Small" not in styles:
        caption_style = styles.add_style("Caption Small", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption_style = styles["Caption Small"]
    set_style_font(caption_style, "Calibri", "微软雅黑", 9, PRESET["muted"])
    caption_style.paragraph_format.space_before = Pt(4)
    caption_style.paragraph_format.space_after = Pt(4)
    caption_style.paragraph_format.line_spacing = 1.1


def configure_list_numbering(doc: Document) -> None:
    numbering = doc.part.numbering_part.element
    abstracts = {
        element.get(qn("w:abstractNumId")): element
        for element in numbering.findall(qn("w:abstractNum"))
    }

    for abstract_id in ("7", "8"):
        abstract = abstracts[abstract_id]
        level = abstract.find(qn("w:lvl"))
        if level is None:
            raise ValueError(f"Missing level 0 for built-in numbering abstract {abstract_id}")

        p_pr = level.find(qn("w:pPr"))
        if p_pr is None:
            p_pr = OxmlElement("w:pPr")
            level.append(p_pr)

        tabs = p_pr.find(qn("w:tabs"))
        if tabs is None:
            tabs = OxmlElement("w:tabs")
            p_pr.append(tabs)
        for child in list(tabs):
            tabs.remove(child)
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)

        ind = p_pr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            p_pr.append(ind)
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")

        spacing = p_pr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            p_pr.append(spacing)
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")


def configure_page(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(PRESET["page_width"])
        section.page_height = Inches(PRESET["page_height"])
        section.top_margin = Inches(PRESET["margin"])
        section.right_margin = Inches(PRESET["margin"])
        section.bottom_margin = Inches(PRESET["margin"])
        section.left_margin = Inches(PRESET["margin"])
        section.header_distance = Inches(PRESET["header_distance"])
        section.footer_distance = Inches(PRESET["footer_distance"])


def paragraph_bottom_border(paragraph, color: str = "D7DBE2", size: int = 6, space: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=PRESET["muted"])
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr_text, fld_sep, fld_text, fld_end])
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=8.5, color=PRESET["muted"])


def configure_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    left = p.add_run("WEATHERPOEM　|　产品需求与实施方案")
    set_run_font(left, size=8.5, color=PRESET["muted"], bold=True)
    paragraph_bottom_border(p, color="E2E5E9", size=4, space=5)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_field(fp)


def add_hyperlink(paragraph, text: str, url: str, color: str = "2E74B5") -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "Hyperlink")
    r_pr.append(r_style)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)"
)


def add_inline_runs(paragraph, text: str, *, size: float | None = None, color: str | None = None) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, ascii_name="Consolas", east_asia="等线", size=(size or 10) - 0.5, color="394556")
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F1F3F5")
            run._element.get_or_add_rPr().append(shd)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color)


def create_numbering(doc: Document, kind: str) -> int:
    numbering = doc.part.numbering_part.element
    if kind == "bullet":
        return 1

    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), "7")
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def add_callout(doc: Document, text: str, label: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.10)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PRESET["callout_fill"])
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "16")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), PRESET["accent"])
    p_bdr.append(left)
    p_pr.append(p_bdr)
    if label:
        run = p.add_run(f"{label}　")
        set_run_font(run, size=10, color=PRESET["accent"], bold=True)
    add_inline_runs(p, text, size=10.2, color="38404A")


def cover_updated_at(markdown_path: Path) -> str:
    traceability_path = (
        markdown_path.resolve().parent.parent
        / "review"
        / "requirements-traceability.json"
    )
    traceability = json.loads(traceability_path.read_text(encoding="utf-8"))
    updated_at = date.fromisoformat(traceability["updatedAt"])
    return f"{updated_at.year} 年 {updated_at.month} 月 {updated_at.day} 日"


def add_cover(doc: Document, updated_at: str) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(86)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(10)
    run = kicker.add_run("PRODUCT REQUIREMENTS & DELIVERY BLUEPRINT")
    set_run_font(run, size=9, color=PRESET["accent"], bold=True)
    run.font.all_caps = True

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("WeatherPoem")
    set_run_font(run, ascii_name="Calibri", east_asia="微软雅黑", size=31, color="18324A", bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(3)
    run = subtitle.add_run("桌面传统文化诗笺")
    set_run_font(run, size=24, color="18324A", bold=True)

    descriptor = doc.add_paragraph()
    descriptor.paragraph_format.space_after = Pt(30)
    run = descriptor.add_run("需求基线、验收门槛与落地实施方案")
    set_run_font(run, size=14, color=PRESET["muted"])

    metadata = [
        ("文档版本", "v3.0"),
        ("文档状态", "内部试用就绪；校园试点待关闭门禁"),
        ("首要用户", "在校大学生"),
        ("首发平台", "Windows 10/11"),
        ("计划验证", "30—50 人，连续 4 周"),
        ("更新日期", updated_at),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    set_table_geometry(table, [1900, 7460], 120)
    set_table_borders(table, color="FFFFFF", size=0)
    set_repeat_table_header(table.rows[0])
    for idx, (label, value) in enumerate(metadata):
        left, right = table.rows[idx].cells
        set_cell_margins(left, top=45, bottom=45, start=120, end=100)
        set_cell_margins(right, top=45, bottom=45, start=120, end=100)
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        rp.paragraph_format.space_after = Pt(0)
        lr = lp.add_run(label)
        set_run_font(lr, size=9.2, color=PRESET["muted"], bold=True)
        rr = rp.add_run(value)
        set_run_font(rr, size=10.2, color=PRESET["ink"])

    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    add_callout(
        doc,
        "让有文化基础但缺少接触契机的大学生，在每天第一次进入电脑桌面时，自然遇见一首与今日情境有关、来源可信且可以继续深入阅读的诗词。",
        "核心价值",
    )

    method = doc.add_paragraph()
    method.paragraph_format.space_before = Pt(20)
    method.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = method.add_run("设计方法　需求 → 决策 → 实现 → 验证 → 风险")
    set_run_font(run, size=9.5, color=PRESET["muted"])
    method.add_run().add_break(WD_BREAK.PAGE)


def add_front_matter(doc: Document) -> None:
    h = doc.add_heading("评审摘要", level=1)
    h.paragraph_format.space_before = Pt(0)
    p = doc.add_paragraph()
    add_inline_runs(
        p,
        "本方案将 WeatherPoem 定义为一个**低打扰、上下文驱动、内容可追溯**的桌面文化产品。"
        "MVP 只验证一个核心命题：大学生是否愿意长期保留一个每天出现一次、不会明显打扰、偶尔能够产生共鸣的传统文化入口。"
        "当前 0.2.0 候选只达到未签名内部受控试用条件，不能据此进入校园试点或公开发布。",
    )

    decisions = [
        "每天每个自然日最多自动展示一次，关闭后只保留托盘入口。",
        "主卡片一次只出现一首诗；文化背景在用户主动展开后呈现。",
        "节日决定主题，天气决定当天感受，节气校正季节意境。",
        "诗词原文、作者和出处来自审核内容库；AI 不创作、不改写原文。",
        "AI 仅在规则选出的最多三首候选中微排序并生成短理由，失败时立即降级。",
        "严重天气预警能够在推荐前阻断普通诗意弹窗。",
        "MVP 采用 Tauri 2 + Vue 3/TypeScript 的本地优先架构。",
        "当前候选以本地规则理由运行；外部 AI、正式预警和公开分发在获得授权与具名批准前保持关闭。",
    ]
    num_id = create_numbering(doc, "bullet")
    for item in decisions:
        lp = doc.add_paragraph()
        apply_numbering(lp, num_id)
        add_inline_runs(lp, item)

    add_callout(
        doc,
        "当前 24 项落地门禁为 7 项通过、16 项部分完成、1 项待执行。内部试用 7/7 就绪；校园试点 7/18、公开发布 7/24，均未就绪。唯一纯待执行项是两名具名内容专家对 120 首作品完成权威底本复核。",
        "评审提醒",
    )

    toc_heading = doc.add_heading("目录", level=1)
    toc_heading.paragraph_format.page_break_before = True
    for chapter_title in CHAPTERS.values():
        entry = doc.add_paragraph()
        entry.paragraph_format.space_after = Pt(5)
        entry.paragraph_format.keep_with_next = True
        run = entry.add_run(chapter_title)
        set_run_font(run, size=10.5, color=PRESET["heading_dark_blue"], bold=True)


def parse_markdown_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    idx = start
    while idx < len(lines) and lines[idx].lstrip().startswith("|"):
        raw_cells = [cell.strip() for cell in lines[idx].strip().strip("|").split("|")]
        rows.append(raw_cells)
        idx += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    return rows, idx


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    rows = [row + [""] * (column_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    set_table_geometry(table, table_widths(column_count, rows), PRESET["table_indent_dxa"])
    set_table_borders(table, PRESET["table_border"], 5)
    set_repeat_table_header(table.rows[0])
    for table_row in table.rows:
        set_row_cant_split(table_row)
    is_version_table = rows[0][:2] == ["版本", "定位"]
    table_font_size = 8.5 if is_version_table else (9.0 if column_count <= 3 else 8.2)

    for row_idx, row in enumerate(rows):
        for col_idx, text in enumerate(row):
            cell = table.rows[row_idx].cells[col_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if is_version_table:
                set_cell_margins(cell, top=45, bottom=45)
            else:
                set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, PRESET["table_fill"])
            elif row_idx % 2 == 0 and len(rows) > 4:
                set_cell_shading(cell, PRESET["table_alt_fill"])
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            if row_idx == 0 or (column_count >= 4 and col_idx == 0):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_runs(
                p,
                text,
                size=table_font_size,
                color=PRESET["ink"],
            )
            if row_idx == 0:
                for run in p.runs:
                    run.bold = True

    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, content: list[str], language: str) -> None:
    if language == "mermaid":
        caption = doc.add_paragraph(style="Caption Small")
        caption_run = caption.add_run("流程/结构示意（文本表达）")
        set_run_font(caption_run, size=9, color=PRESET["muted"], bold=True)
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F5F7")
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "5")
        el.set(qn("w:color"), "E0E4E8")
        p_bdr.append(el)
    p_pr.append(p_bdr)
    for idx, line in enumerate(content):
        run = p.add_run(line)
        set_run_font(run, ascii_name="Consolas", east_asia="等线", size=8.2, color="394556")
        if idx < len(content) - 1:
            run.add_break()


def section_number(text: str) -> int | None:
    match = re.match(r"##\s+(\d+)\.", text)
    return int(match.group(1)) if match else None


def add_document_body(doc: Document, markdown_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    start_idx = next(i for i, line in enumerate(lines) if line.startswith("## 1."))
    lines = lines[start_idx:]

    idx = 0
    skip_section = False
    in_code = False
    code_language = ""
    code_lines: list[str] = []
    current_list_kind: str | None = None
    current_num_id: int | None = None

    while idx < len(lines):
        line = lines[idx].rstrip()

        if in_code:
            if line.startswith("```"):
                add_code_block(doc, code_lines, code_language)
                in_code = False
                code_language = ""
                code_lines = []
            else:
                code_lines.append(line)
            idx += 1
            continue

        if line.startswith("```"):
            in_code = True
            code_language = line[3:].strip()
            code_lines = []
            idx += 1
            continue

        if line.startswith("## "):
            sec_no = section_number(line)
            skip_section = sec_no in SKIP_SECTIONS if sec_no is not None else False
            if skip_section:
                idx += 1
                continue
            if sec_no in CHAPTERS:
                chapter = doc.add_heading(CHAPTERS[sec_no], level=1)
                if sec_no != 1:
                    chapter.paragraph_format.page_break_before = True
            heading_text = re.sub(r"^##\s+", "", line)
            doc.add_heading(heading_text, level=2)
            current_list_kind = None
            current_num_id = None
            idx += 1
            continue

        if skip_section:
            idx += 1
            continue

        if line.startswith("### "):
            heading_text = re.sub(r"^###\s+", "", line)
            doc.add_heading(heading_text, level=3)
            current_list_kind = None
            current_num_id = None
            idx += 1
            continue

        if line.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(re.sub(r"^####\s+", "", line))
            set_run_font(run, size=10.5, color=PRESET["heading_dark_blue"], bold=True)
            current_list_kind = None
            current_num_id = None
            idx += 1
            continue

        if line.lstrip().startswith("|"):
            rows, next_idx = parse_markdown_table(lines, idx)
            add_table(doc, rows)
            current_list_kind = None
            current_num_id = None
            idx = next_idx
            continue

        if line.startswith("> "):
            add_callout(doc, line[2:].strip())
            current_list_kind = None
            current_num_id = None
            idx += 1
            continue

        bullet_match = re.match(r"^\s*-\s+(.+)", line)
        number_match = re.match(r"^\s*\d+\.\s+(.+)", line)
        if bullet_match or number_match:
            kind = "bullet" if bullet_match else "decimal"
            if current_list_kind != kind or current_num_id is None:
                current_num_id = create_numbering(doc, kind)
                current_list_kind = kind
            p = doc.add_paragraph()
            apply_numbering(p, current_num_id)
            p.paragraph_format.space_after = Pt(4)
            add_inline_runs(p, (bullet_match or number_match).group(1))
            idx += 1
            continue

        if not line.strip():
            current_list_kind = None
            current_num_id = None
            idx += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.keep_together = False
        add_inline_runs(p, line.strip())
        current_list_kind = None
        current_num_id = None
        idx += 1

    if in_code and code_lines:
        add_code_block(doc, code_lines, code_language)


def set_document_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "WeatherPoem 桌面传统文化诗笺 - 需求基线、验收门槛与落地实施方案"
    props.subject = "传统文化桌面产品的需求基线、交互、AI、架构、实施、试点与发布验收"
    props.author = "WeatherPoem 项目组"
    props.keywords = "WeatherPoem, 传统文化, 诗词, 桌面应用, PRD, AI, MVP"
    props.comments = "v3.0 内部试用就绪、校园试点与公开发布待关闭门禁"


def build_document(markdown_path: Path, output_path: Path) -> None:
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    configure_list_numbering(doc)
    configure_header_footer(doc)
    set_document_properties(doc)
    add_cover(doc, cover_updated_at(markdown_path))
    add_front_matter(doc)
    add_document_body(doc, markdown_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build_document(args.markdown, args.output)
    print(f"Created {args.output}")


if __name__ == "__main__":
    main()
