from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "output" / "WeatherPoem_项目需求与落地实施方案_v3.0.docx"
OUTPUT = ROOT / "output" / "WeatherPoem_项目需求与落地实施方案_v3.1.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_status_table(document: Document, rows: list[tuple[str, str, str, str]]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ("范围", "当前结论", "已有证据", "仍需完成")
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, "24413A")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index])
            if row_index % 2 == 0:
                set_cell_shading(cells[index], "F3F0E8")
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    widths = (Cm(2.6), Cm(2.8), Cm(6.2), Cm(6.2))
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = width


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(text)


def main() -> None:
    manifest = json.loads(
        (ROOT / "output" / "WeatherPoem_0.2.0_RELEASE.json").read_text(encoding="utf-8")
    )
    readiness = json.loads(
        (ROOT / "output" / "WeatherPoem_0.2.0_READINESS.json").read_text(encoding="utf-8")
    )
    artifacts = {item["kind"]: item for item in manifest["artifacts"]}

    document = Document(SOURCE)
    document.add_section(WD_SECTION.NEW_PAGE)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("附录 A｜内部试用候选收尾记录")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(36, 65, 58)

    subtitle = document.add_paragraph("版本 3.1 · 2026-09-10 · 与当前 0.2.0 未签名试用候选绑定")
    subtitle.paragraph_format.space_after = Pt(14)
    for run in subtitle.runs:
        run.italic = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(92, 101, 96)

    document.add_heading("A.1 结论与适用边界", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("当前结论：").bold = True
    paragraph.add_run(
        "WeatherPoem 0.2.0 达到“未签名、受控、小范围内部试用”条件；"
        "不等于 30—50 人校园试点就绪，也不等于公开发布就绪。"
    )
    paragraph = document.add_paragraph()
    paragraph.add_run("本轮决策：").bold = True
    paragraph.add_run(
        "按项目方决定跳过七天正式长测。七天跨日、休眠唤醒与真实在线链路仍为未完成项，"
        "RQT-13 继续保持 partial；旧候选的八小时报告仅作归档，不绑定本候选。"
    )

    milestones = {item["id"]: item for item in readiness["milestones"]}
    add_status_table(
        document,
        [
            (
                "内部试用",
                "READY",
                "核心回归、性能冒烟、受控故障矩阵、短时资源演练、NSIS 生命周期、MSI 镜像、单实例与窗口定位均通过。",
                "仅限知情、受控分发；保持未签名提示与人工反馈通道。",
            ),
            (
                "校园试点",
                "NOT READY",
                f"{milestones['campus-pilot']['passedGateCount']}/{milestones['campus-pilot']['requiredGateCount']} 项门禁通过。",
                "双人内容复核、真实设备矩阵、管理员 MSI/重启、自启、七天长稳、可访问性人工验证、UAT 与具名批准。",
            ),
            (
                "公开发布",
                "NOT READY",
                f"{milestones['public-release']['passedGateCount']}/{milestones['public-release']['requiredGateCount']} 项门禁通过。",
                "在校园试点门禁之外，还需代码签名、签名更新、公开分发与撤回演练、许可和运维责任人签署。",
            ),
        ],
    )

    document.add_heading("A.2 本轮关键修复与验证", level=1)
    add_bullet(
        document,
        "修复 Tauri 隐藏窗口仍默认抢占焦点的问题：主窗口改为 visible=false、focus=false，避免新版 WebView2 初始化后只剩托盘进程。",
    )
    add_bullet(
        document,
        "Vite 产物改为相对资源路径，确保桌面本地协议从 ./assets/ 加载脚本与样式，并用契约测试锁定。",
    )
    add_bullet(
        document,
        "修复 Node 子进程启动 PowerShell 时模块无法自动加载的问题；签名与哈希脚本显式加载系统模块。",
    )
    add_bullet(
        document,
        "npm 已知漏洞由 5 项降至 0；Rust event-listener 从 5.4.1 升至 5.4.2，Windows 目标 RustSec 漏洞与 unsound 阻断项均为 0。",
    )
    add_bullet(
        document,
        "全量 npm run check 与 release:verify 已通过；第三方 NOTICE 已实际进入当前 NSIS/MSI，自动结论仍不替代具名法务批准。",
    )

    document.add_heading("A.3 当前候选身份", level=1)
    identity = document.add_table(rows=1, cols=3)
    identity.style = "Table Grid"
    identity.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, value in enumerate(("发布物", "SHA-256", "可信状态")):
        identity.rows[0].cells[index].text = value
        set_cell_shading(identity.rows[0].cells[index], "24413A")
        for run in identity.rows[0].cells[index].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for kind in ("portable", "nsis", "msi"):
        item = artifacts[kind]
        cells = identity.add_row().cells
        cells[0].text = item["name"]
        cells[1].text = item["sha256"]
        cells[2].text = "NotSigned / unsigned-trial"
        for cell in cells:
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(7.5 if cell is cells[1] else 8.5)

    document.add_heading("A.4 下一阶段的执行顺序", level=1)
    for text in (
        "项目发起人先指定内容、质量、Windows 发布、安全隐私、试点、法务和运维的真实责任人，并冻结内部试用范围。",
        "由两名不同内容专家逐首复核 120 首诗词的权威底本、作者、篇名、断句、译注与授权证据。",
        "在专用 Windows 10/11 账户完成管理员 MSI、两次重启、自启启停、125%—200% 缩放、多显示器和屏幕阅读器人工验收。",
        "如项目方继续跳过七天长测，则校园试点保持 not-ready；若恢复该门禁，应冻结候选并在持续供电、禁止睡眠的专用主机上重新执行。",
        "两名未参与开发的大学生完成 UAT 后，再召开校园试点 Go/No-Go；不得直接进入公开发布。",
    ):
        add_bullet(document, text)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(0)
    run = note.add_run(
        "审计口径：所有自动证据只证明对应脚本、候选字节与受控场景；"
        "真实用户效果、法律授权、组织责任与公开分发资格必须由具名人员完成。"
    )
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(92, 101, 96)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
