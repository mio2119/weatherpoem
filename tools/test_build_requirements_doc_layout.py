from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


TOOLS_DIR = Path(__file__).resolve().parent
if str(TOOLS_DIR) not in sys.path:
    sys.path.insert(0, str(TOOLS_DIR))

import build_requirements_doc as builder


class RequirementsDocLayoutTests(unittest.TestCase):
    def test_cover_uses_traceability_updated_at(self) -> None:
        with tempfile.TemporaryDirectory() as temp_directory:
            project_root = Path(temp_directory)
            docs_directory = project_root / "docs"
            review_directory = project_root / "review"
            docs_directory.mkdir()
            review_directory.mkdir()
            markdown_path = docs_directory / "requirements.md"
            markdown_path.write_text(
                "# WeatherPoem\n\n## 1. 执行摘要\n\n正文。\n",
                encoding="utf-8",
            )
            (review_directory / "requirements-traceability.json").write_text(
                json.dumps({"updatedAt": "2026-08-05"}),
                encoding="utf-8",
            )
            output_path = project_root / "requirements.docx"

            builder.build_document(markdown_path, output_path)

            generated = Document(output_path)
            metadata_pairs = {
                (row.cells[0].text, row.cells[1].text)
                for table in generated.tables
                for row in table.rows
                if len(row.cells) >= 2
            }
            self.assertIn(
                ("更新日期", "2026 年 8 月 5 日"),
                metadata_pairs,
                "The cover date must follow the authoritative traceability metadata.",
            )

    def test_code_block_keeps_all_lines_on_one_page(self) -> None:
        doc = Document()
        builder.configure_styles(doc)

        builder.add_code_block(doc, ["first", "second", "third"], "json")

        code_paragraph = doc.paragraphs[-1]
        keep_lines = code_paragraph._p.pPr.find(qn("w:keepLines"))
        self.assertIsNotNone(
            keep_lines,
            "Code blocks must not leave a few trailing lines orphaned on the next page.",
        )

    def test_state_machine_table_reserves_width_for_identifiers(self) -> None:
        rows = [
            ["状态", "含义", "允许进入", "退出条件"],
            ["NOT_CONFIGURED", "首次设置未完成", "首次启动", "设置保存"],
            ["ALERT_SUPPRESSED", "严重预警暂停", "PREPARING", "新自然日"],
            ["RETRYABLE_FAILURE", "可恢复故障", "RENDERING", "下次启动"],
        ]

        widths = builder.table_widths(4, rows)

        self.assertEqual(sum(widths), builder.PRESET["content_width_dxa"])
        self.assertGreaterEqual(
            widths[0],
            2200,
            "State identifiers must not be split in the middle of a token.",
        )

    def test_three_column_field_table_reserves_width_for_identifiers(self) -> None:
        rows = [
            ["字段", "必填", "说明"],
            ["evidence_checked_at", "是", "最近核对时间"],
            ["cultural_background", "是", "文化背景"],
            ["contemporary_connection", "是", "与大学生活的联系"],
        ]

        widths = builder.table_widths(3, rows)

        self.assertEqual(sum(widths), builder.PRESET["content_width_dxa"])
        self.assertGreaterEqual(
            widths[0],
            2300,
            "Long field identifiers must wrap at underscores, not in the middle of a word.",
        )


if __name__ == "__main__":
    unittest.main()
